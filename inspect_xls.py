import openpyxl
import json

wb1 = openpyxl.load_workbook("PROCESSOS CCM 02_10.xlsx", data_only=True)
ws1 = wb1.active

processos_data = []
for row in ws1.iter_rows(values_only=True):
    if any(row):
        processos_data.append(row)

print("PROCESSOS:", json.dumps(processos_data[:20], indent=2, ensure_ascii=False))
print(f"Total rows: {len(processos_data)}")

wb2 = openpyxl.load_workbook("MAPEAMENTO - CCM (1).xlsx", data_only=True)
ws2 = wb2['ATENDIMENTOS']

atendimentos_data = []
for row in ws2.iter_rows(values_only=True):
    if any(row):
        atendimentos_data.append(row)

print("\nATENDIMENTOS:", json.dumps(atendimentos_data, indent=2, ensure_ascii=False))

