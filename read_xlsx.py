# -*- coding: utf-8 -*-
import openpyxl, json, os, sys

# Listar arquivos xlsx no diretório
xlsx_files = [f for f in os.listdir('.') if f.endswith('.xlsx')]
print("Arquivos XLSX encontrados:", xlsx_files)

if not xlsx_files:
    print("Nenhum XLSX encontrado!")
    sys.exit(1)

filename = xlsx_files[0]
print(f"\nLendo: {filename}\n")

wb = openpyxl.load_workbook(filename, data_only=True)
ws = wb.active

# Cabecalho
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print('TOTAL COLUNAS:', len(headers))
print()
for i, h in enumerate(headers):
    print(f'Col {i:02d}: {repr(h)}')

print('\n\n======= TOTAL DE LINHAS =======')
all_rows = list(ws.iter_rows(min_row=2, values_only=True))
linhas_com_dados = [r for r in all_rows if any(v is not None for v in r)]
print(f'Respostas: {len(linhas_com_dados)}')

print('\n\n======= TODAS AS RESPOSTAS (JSON) =======')
result = []
for row in linhas_com_dados:
    entry = {}
    for i, val in enumerate(row):
        if val is not None and i < len(headers) and headers[i]:
            key = str(headers[i]).strip()
            if hasattr(val, 'isoformat'):
                entry[key] = val.isoformat()
            else:
                entry[key] = val
    result.append(entry)

with open('dados_cartografia.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

print("JSON salvo em dados_cartografia.json")
print(f"Total de respostas processadas: {len(result)}")
