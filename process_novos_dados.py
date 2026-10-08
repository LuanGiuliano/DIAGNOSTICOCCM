import openpyxl
import json
import re

def clean_text(text):
    if not isinstance(text, str): return text
    # basic encoding fix if needed, but openpyxl usually returns unicode.
    # The terminal output had '' because of terminal encoding, openpyxl should have actual characters.
    return text.strip()

processos = []
wb1 = openpyxl.load_workbook("PROCESSOS CCM 02_10.xlsx", data_only=True)
ws1 = wb1.active
for row in ws1.iter_rows(values_only=True):
    caixa = row[0]
    if isinstance(caixa, str) and 'CCM' in caixa:
        pae3 = row[1] if row[1] is not None else 0
        pae4 = row[2] if row[2] is not None else 0
        
        # extrair nome simplificado da caixa
        nome_curto = caixa
        if 'TRI' in caixa: nome_curto = 'TRIÊNIO'
        elif 'CERTID' in caixa: nome_curto = 'CERTIDÕES'
        elif 'DESIG' in caixa: nome_curto = 'DESIGNAÇÃO'
        elif 'FÉRIAS' in caixa or 'FERIAS' in caixa: nome_curto = 'FÉRIAS'
        elif 'LICEN' in caixa: nome_curto = 'LICENÇAS'
        elif 'PECÚNIA' in caixa or 'PECUNIA' in caixa: nome_curto = 'PECÚNIA'
        elif 'PROTOCOLO' in caixa: nome_curto = 'PROTOCOLO'
        elif 'DIÁRIO' in caixa or 'DIARIO' in caixa: nome_curto = 'DIÁRIO'
        elif 'AVERB' in caixa: nome_curto = 'AVERBAÇÃO'
        elif 'COMUNICA' in caixa: nome_curto = 'COMUNICAÇÃO'
        elif 'ARQUIVO' in caixa: nome_curto = 'ARQUIVO'
        elif 'DECLARA' in caixa: nome_curto = 'DECLARAÇÃO'
        elif 'MANUTEN' in caixa: nome_curto = 'MANUTENÇÃO'
        elif 'POLO' in caixa: nome_curto = 'POLO'
        elif 'CCM- COORD' in caixa: nome_curto = 'CCM (Geral)'
        else: nome_curto = caixa.split('>')[1].strip() if '>' in caixa else caixa

        if (pae3 > 0 or pae4 > 0):
            processos.append({
                "caixa": nome_curto,
                "pae3": pae3,
                "pae4": pae4,
                "total": pae3 + pae4
            })

atendimentos = []
wb2 = openpyxl.load_workbook("MAPEAMENTO - CCM (1).xlsx", data_only=True)
ws2 = wb2['ATENDIMENTOS']
for idx, row in enumerate(ws2.iter_rows(values_only=True)):
    if idx == 0: continue # skip header
    carteira = row[2]
    if carteira and carteira != 'TOTAL':
        servidores = row[3] if row[3] is not None else 0
        estagiarios = row[4] if row[4] is not None else 0
        atendimentos.append({
            "carteira": clean_text(carteira),
            "servidores": servidores,
            "estagiarios": estagiarios,
            "total": servidores + estagiarios
        })

output = {
    "processos_caixas": processos,
    "atendimentos_carteira": atendimentos
}

with open("src/data/novos_dados_ccm.json", "w", encoding="utf-8") as f:
    json.dump(output, f, ensure_ascii=False, indent=2)

print("Gerado src/data/novos_dados_ccm.json com sucesso!")
